import asyncio
import os
import json
import pathlib
from datetime import datetime
from dateutil import parser as date_parser
from dotenv import load_dotenv
import google.generativeai as genai
from docx import Document
import re
from typing import Dict, List, Tuple, Any
from collections import defaultdict
from app.db.database import get_database
import random
import uuid
from fastapi import HTTPException
from bson import ObjectId
from app.utils.prompts import (
    CV_EXTRACTION_PROMPT,
    MISSING_FIELD_SUGGESTIONS_PROMPT,
)


load_dotenv()
genai.configure(api_key=os.getenv("GEMINI_API_KEY"))
model = genai.GenerativeModel("gemini-2.5-flash-lite")



async def _gemini_with_retry(prompt, max_retries: int = 4):
    """Call Gemini with exponential backoff on rate-limit (429) errors."""
    delay = 2.0
    for attempt in range(max_retries):
        try:
            return await model.generate_content_async(prompt)
        except Exception as exc:
            err = str(exc)
            is_rate_limit = "429" in err or "ResourceExhausted" in err or "quota" in err.lower()
            if is_rate_limit and attempt < max_retries - 1:
                await asyncio.sleep(delay)
                delay *= 2
                continue
            raise

def parse_duration(duration_str):
    try:
        # --- Normalize input ---
        duration_str = duration_str.replace("'", "'").replace("'", "'")
        duration_str = re.sub(r"[–—−]", "-", duration_str).strip()
        duration_str = re.sub(r"\s+", " ", duration_str)
        duration_str = re.sub(r"\(.*?\)", "", duration_str)  # remove (...) notes

        # --- Handle "Since <date>" ---
        if duration_str.lower().startswith("since"):
            start_str = duration_str[5:].strip()
            try:
                start_date = date_parser.parse(start_str, fuzzy=True)
            except:
                return None
            return start_date, datetime.today()

        # --- Split into start/end parts (accepts -, to, –) ---
        parts = re.split(r"\s*(?:-|to)\s*", duration_str, flags=re.IGNORECASE)
        if len(parts) != 2:
            return None

        start_str, end_str = parts[0].strip(), parts[1].strip().lower()

        # --- Parse start date ---
        start_date = None
        for fmt in ["%m/%Y", "%b'%y", "%b %y"]:
            try:
                start_date = datetime.strptime(start_str, fmt)
                break
            except:
                continue
        if not start_date:
            start_date = date_parser.parse(start_str, fuzzy=True)

        # --- Parse end date ---
        if any(word in end_str for word in ["present", "current", "now"]):
            end_date = datetime.today()
        else:
            end_date = None
            for fmt in ["%m/%Y", "%b'%y", "%b %y"]:
                try:
                    end_date = datetime.strptime(end_str, fmt)
                    break
                except:
                    continue
            if not end_date:
                end_date = date_parser.parse(end_str, fuzzy=True)

        return start_date, end_date

    except Exception as e:
        print("Duration parsing error:", str(e))
        return None
def extract_years_from_summary(summary_text: str) -> float:
    """
    Extract years of experience directly from the Summary section
    if WorkExperience durations are missing.
    """
    if not summary_text:
        return 0.0

    match = re.search(r"(?i)(over|more than|about)?\s*(\d+)\s*(\+)?\s*years? of experience", summary_text)
    if match:
        return float(match.group(2))
    return 0.0



def calculate_years_of_experience(work_experiences):
    total_months = 0
    for job in work_experiences:
        duration_str = job.get("Duration", "")
        parsed = parse_duration(duration_str)
        if parsed:
            start, end = parsed
            months = (end.year - start.year) * 12 + (end.month - start.month) + 1
            total_months += max(0, months)
    return round(total_months / 12, 2)


def clean_json(obj):
    if isinstance(obj, dict):
        return {k: clean_json(v) for k, v in obj.items()}
    elif isinstance(obj, list):
        return [clean_json(v) for v in obj]
    elif obj is None:
        return ""
    return obj


def extract_docx_text(filepath):
    doc = Document(filepath)
    texts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                cell_text = cell.text.strip()
                if cell_text:
                    texts.append(cell_text)
    return "\n".join(texts)
def extract_job_role(parsed_json):

    work_exps = parsed_json.get("WorkExperience", [])
    if work_exps and work_exps[0].get("Role"):
        role = work_exps[0]["Role"]

        return role.split("/")[0].split(",")[0].strip()

    summary = parsed_json.get("Summary", "")
    if summary:
        match = re.search(r"(?i)([A-Z][a-zA-Z\s\/\-]+)\s+with\s+\d+\s+years", summary)
        if match:
            role = match.group(1).strip()
            return role.split("/")[0].split(",")[0].strip()

    return None




async def extract_cv_data_from_file(filepath: str, mime_type: str):
    
    prompt = CV_EXTRACTION_PROMPT

    try:
        if mime_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            file_text = extract_docx_text(filepath)
            content_input = [file_text, prompt]
        else:
            file_bytes = pathlib.Path(filepath).read_bytes()
            content_input = [
                {"mime_type": mime_type, "data": file_bytes},
                prompt
            ]

        response = await _gemini_with_retry(content_input)
        response_text = response.text.strip()

        if response_text.startswith("```json"):
            response_text = response_text[7:]
        if response_text.endswith("```"):
            response_text = response_text[:-3]

        response_text = response_text.strip()
        start_idx = response_text.find('{')
        end_idx = response_text.rfind('}')
        response_text = response_text[start_idx:end_idx+1]

        parsed_json = json.loads(response_text)
        parsed_json = clean_json(parsed_json)

        work_exp = parsed_json.get("WorkExperience", [])
        years_from_work = calculate_years_of_experience(work_exp)

        if years_from_work > 0:
            parsed_json["YearsOfExperience"] = years_from_work
        else:
            parsed_json["YearsOfExperience"] = extract_years_from_summary(parsed_json.get("Summary", ""))

        parsed_json["JobRole"] = extract_job_role(parsed_json)

        return parsed_json

    except Exception as e:
        print("CV Extraction Error:", str(e))
        return {"error": "Failed to parse CV data"}


def predict_audience_type(parsed_data: Dict) -> str:
    work_exp_list = parsed_data.get("WorkExperience", [])

    work_periods = []
    for job in work_exp_list:
        duration_str = job.get("Duration", "")
        parsed = parse_duration(duration_str)
        if parsed:
            work_periods.append(parsed)

    work_periods.sort(key=lambda x: x[0])

    merged_work: List[Tuple[datetime, datetime]] = []
    for p in work_periods:
        if not merged_work:
            merged_work.append(p)
        else:
            last_start, last_end = merged_work[-1]
            curr_start, curr_end = p
            if curr_start <= last_end:  
                merged_work[-1] = (last_start, max(last_end, curr_end))
            else:
                merged_work.append(p)

    if isinstance(parsed_data.get("YearsOfExperience"), (int, float)):
        total_years = float(parsed_data["YearsOfExperience"])
    else:
        total_years = sum((e - s).days for s, e in merged_work) / 365.0 if merged_work else 0.0

    currently_working = any(
        ("Duration" in w and re.search(r"(present|current|ongoing|now)", w["Duration"], re.IGNORECASE))
        for w in work_exp_list
    )

    def months_between(a: datetime, b: datetime) -> float:
        return (b - a).days / 30.0

    employment_gap = False
    if not currently_working and merged_work:
        last_end = merged_work[-1][1]
        if months_between(last_end, datetime.today()) >= 12:
            employment_gap = True

    if not currently_working and employment_gap:
        return "Job Seeker"
    elif total_years == 0:
        return "Student"
    elif 0 < total_years <= 5:
        return "Early Professional (2-3 years of experience)"
    elif total_years > 5:
        return "Mid Career Pivot"
    else:
        return "Early Professional (2-3 years of experience)"

FIELD_MAPPING = {
    "softskills": "softskills_suggestions",
    "soft skills": "softskills_suggestions",
    "softskills_suggestions": "softskills_suggestions",

    "hardskills": "technical_skills_suggestions",
    "hard skills": "technical_skills_suggestions",
    "tools": "technical_skills_suggestions",
    "technicalskills": "technical_skills_suggestions",
    "technical_skills": "technical_skills_suggestions",

    "certifications": "certifications_suggestions",
    "certification": "certifications_suggestions",
    "certificate": "certifications_suggestions",
    "certificates": "certifications_suggestions",
    "certification_suggestions": "certifications_suggestions"
}

def normalize_key(key: str) -> str:
    return key.strip().lower().replace("_", "").replace(" ", "")

def clean_llm_json_response(response_text: str) -> str:
    response_text = response_text.strip()
    if response_text.startswith("```json"):
        response_text = response_text[len("```json"):].strip()
    if response_text.startswith("```"):
        response_text = response_text[3:].strip()
    if response_text.endswith("```"):
        response_text = response_text[:-3].strip()

    start_idx = response_text.find("{")
    end_idx = response_text.rfind("}")
    if start_idx != -1 and end_idx != -1:
        response_text = response_text[start_idx:end_idx + 1]

    return response_text


async def generate_missing_field_suggestions(cv_context: dict) -> dict:
    suggestions = {
        "softskills_suggestions": [],
        "technical_skills_suggestions": [],
        "certifications_suggestions": []
    }

    skills = cv_context.get("Skills", {}) or {}
    missing_fields = []
    if not skills.get("SoftSkills"):
        missing_fields.append("SoftSkills")
    if not skills.get("HardSkills"):
        missing_fields.append("HardSkills")
    if not cv_context.get("Certifications"):
        missing_fields.append("Certifications")

    if not missing_fields:
        return suggestions

    context_str = "\n".join(f"{k}: {v}" for k, v in cv_context.items() if v)
    prompt = MISSING_FIELD_SUGGESTIONS_PROMPT.format(
        context_str=context_str,
        missing_fields=missing_fields,
    )

    response = await _gemini_with_retry(prompt)
    cleaned = clean_llm_json_response(response.text)

    try:
        parsed = json.loads(cleaned)
    except Exception:
        parsed = {f: [] for f in missing_fields}
    for key, value in parsed.items():
        norm_key = normalize_key(key)
        mapped_field = FIELD_MAPPING.get(norm_key)

        if not mapped_field:
            continue

        if isinstance(value, str):
            items = [s.strip() for s in value.split(",") if s.strip()]
        elif isinstance(value, list):
            items = [str(s).strip() for s in value if str(s).strip()]
        else:
            items = []

        if mapped_field == "softskills_suggestions" and "SoftSkills" in missing_fields:
            suggestions["softskills_suggestions"].extend(items)

        if mapped_field == "technical_skills_suggestions" and "HardSkills" in missing_fields:
            suggestions["technical_skills_suggestions"].extend(items)

        if mapped_field == "certifications_suggestions" and "Certifications" in missing_fields:
            suggestions["certifications_suggestions"].extend(items)

    for k in suggestions:
        suggestions[k] = list(set(suggestions[k]))


    return suggestions


async def generate_job_attribute_options(cv_context: dict, questions_from_db: list) -> dict:
    """
    Generates career-related multiple-choice options for each question.

    Rules:
    - If single parameter → 5 options.
    - If multiple parameters joined with '+' → 2 options per parameter.
    - Considers both parameters and question meaning to generate options.
    - Outputs a single list of options.
    - Each option starts with a capital letter.
    """

    # ---- Extract the required fields ----
    summary = cv_context.get("Summary")
    experience = cv_context.get("Experience") or cv_context.get("WorkExperience")
    industry = cv_context.get("Industry") or cv_context.get("IndustryDomain")
    domain = cv_context.get("Domain")
    technical_skills = cv_context.get("TechnicalSkills") or cv_context.get("Skills", {}).get("HardSkills")
    soft_skills = cv_context.get("SoftSkills") or cv_context.get("Skills", {}).get("SoftSkills")
    talent_attributes = cv_context.get("TalentAttributes") or cv_context.get("Talent Information")

    # ---- DEDUPLICATION HELPER ----
    def dedupe(value):
        if isinstance(value, list):
            seen = set()
            unique = []
            for v in value:
                key = str(v)
                if key not in seen:
                    seen.add(key)
                    unique.append(v)
            return unique

        if isinstance(value, str):
            # Remove duplicate lines
            lines = value.split("\n")
            unique_lines = list(dict.fromkeys(lines))
            return "\n".join(unique_lines)

        return value  # dict or others

    # ---- Build CLEAN context with NO duplicates ----
    filtered_context = {
        "Summary": dedupe(summary),
        "Experience": dedupe(experience),
        "Industry": dedupe(industry),
        "Domain": dedupe(domain),
        "TechnicalSkills": dedupe(technical_skills),
        "SoftSkills": dedupe(soft_skills),
        "TalentAttributes": dedupe(talent_attributes),
    }

    # ---- Convert filtered items into a string for LLM ----
    context_str = "\n".join(
        f"{k}: {v}" for k, v in filtered_context.items() if v
    )

    # ---- Audience filter ----
    audience_type = cv_context.get("audience_type") or cv_context.get("audience")

    if audience_type:
        questions_from_db = [
            q for q in questions_from_db
            if not q.get("audience") or q.get("audience") == audience_type
        ]

    questions_meta = []
    for q in questions_from_db:
        parameter = q.get("parameter", "")
        parameter_list = [p.strip() for p in parameter.split(" + ")]
        option_count = 5 if len(parameter_list) == 1 else 2 * len(parameter_list)
        questions_meta.append({
            "parameter": parameter,
            "parameter_list": parameter_list,
            "question": q.get("question"),
            "type": q.get("type"),
            "iconfilename": q.get("iconfilename"),
            "limit": q.get("limit") or q.get("Limit"),
            "option_count": option_count,
        })

    batch_prompt = (
        "Generate focused, high-quality multiple-choice options for each question below, "
        "based on the user's professional background and the intent of each question.\n\n"
        "Generate options that introduce new elements not already in the resume/CV but relatable to the job title.\n\n"
        "CONTEXT SUMMARY:\n"
        f"{context_str}\n\n"
        "QUESTIONS (JSON array):\n"
        f"{json.dumps([{'parameter': m['parameter'], 'question': m['question'], 'option_count': m['option_count']} for m in questions_meta], indent=2)}\n\n"
        "REQUIREMENTS:\n"
        "- For each question, provide EXACTLY the number of options given in option_count.\n"
        "- Options must be short (2-5 words), clear, and directly related to the parameter.\n"
        "- Each option must start with a capital letter.\n"
        "- No numbering, bullets, special symbols, or prefixes.\n"
        "- Avoid generic, vague, or repetitive wording.\n\n"
        "Return ONLY valid JSON:\n"
        "{\"results\": [{\"parameter\": \"<exact parameter string>\", \"options\": [\"Option 1\", ...]}, ...]}\n"
    )

    suggestions = []
    try:
        response = await _gemini_with_retry(batch_prompt)
        cleaned = clean_llm_json_response(response.text)
        parsed_batch = json.loads(cleaned)
        results_map = {r["parameter"]: r.get("options", []) for r in parsed_batch.get("results", [])}

        for m in questions_meta:
            options = results_map.get(m["parameter"], [])
            formatted = []
            for opt in options[:m["option_count"]]:
                opt = re.sub(r"^[^A-Za-z]+", "", opt.strip())
                formatted.append((opt[:1].upper() + opt[1:]) if opt else "Option")
            while len(formatted) < m["option_count"]:
                formatted.append(f"Option {len(formatted) + 1}")

            item = {
                "parameter": m["parameter_list"],
                "question": m["question"],
                "type": m["type"],
                "iconfilename": m["iconfilename"],
                "options": formatted,
            }
            if m["limit"] is not None:
                item["limit"] = m["limit"]
            suggestions.append(item)
    except Exception:
        for m in questions_meta:
            item = {
                "parameter": m["parameter_list"],
                "question": m["question"],
                "type": m["type"],
                "iconfilename": m["iconfilename"],
                "options": [f"Option {i + 1}" for i in range(m["option_count"])],
            }
            if m["limit"] is not None:
                item["limit"] = m["limit"]
            suggestions.append(item)

    return {
        "success": True,
        "suggestions": suggestions,
    }

async def get_latest_user_cv(db, user_id: str):
    """
    Fetch the most recent uploaded CV for a logged-in user.
    """
    latest_cv = await db["uploads"].find_one(
        {"user_id": ObjectId(user_id), "source": "cv"},
        sort=[("uploaded_at", -1)]
    )
    return latest_cv


async def generate_anchor_attribute_options(user_id: str, questions, model, get_database):
    """
    Generate multiple-choice options for Anchor attributes based on:
    - Personal Interests + Hobbies + Exploration Interest + Motivation Drivers + Motivating Activities
    - Achievements
    - PLUS resume context: Summary + Experience text + Education text

    Each parameter contributes EXACTLY 2 options.
    Saves results in the latest uploaded CV document for the user.
    """

    target_parameters = {
        "Creative Inclinations + Organizational Skills + Competency + Personality Traits",
        "Newly Acquired Skills + Emerging Tech Awareness + Future Study Intent"
    }

    db = await get_database()

    # ---- Get latest CV upload ----
    latest_cv = await get_latest_user_cv(db, user_id)
    if not latest_cv or "parsed_data" not in latest_cv:
        raise HTTPException(status_code=404, detail="No CV found for this user")

    cv_id = str(latest_cv["_id"])

    # ---- Fetch only required anchor answers ----
    cursor = db["answers"].find({
        "cv_id": cv_id,
        "section": "Anchor Attributes",
        "parameter": {
            "$in": [
                "Personal Interests + Hobbies + Exploration Interest + Motivation Drivers + Motivating Activities",
                "Achievements"
            ]
        }
    })

    answers = await cursor.to_list(length=None)
    if not answers:
        raise HTTPException(status_code=404, detail="No required anchor answers found")

    # ---- Extract free text for anchor parameters ----
    base_free_text_map = {}
    for ans in answers:
        param = ans["parameter"]
        val = ans["value"]
        if isinstance(val, str):
            base_free_text_map[param] = val
        elif isinstance(val, dict) and "text" in val:
            base_free_text_map[param] = val["text"]

    non_empty_texts = [t for t in base_free_text_map.values() if t]
    random.shuffle(non_empty_texts)
    context_sample = "\n".join(non_empty_texts)

    # ----------------------------------------------------------
    # 📌 ADD RESUME TEXT FROM UPLOADS → Summary, Experience, Education
    # ----------------------------------------------------------
    parsed = latest_cv.get("parsed_data", {})

    # Summary
    summary_text = parsed.get("Summary") or ""

    # Experience text (joining all descriptions)
    experience_text = ""
    work_exp = parsed.get("WorkExperience") or []
    if isinstance(work_exp, list):
        experience_text = "\n".join([
            w.get("Description", "") for w in work_exp if w.get("Description")
        ])

    # Education text
    education_text = ""
    edu_list = parsed.get("Education") or []
    if isinstance(edu_list, list):
        education_text = "\n".join([
            f"{e.get('Degree', '')} {e.get('College', '')} {e.get('Year', '')}".strip()
            for e in edu_list
            if e.get('Degree') or e.get('College') or e.get('Year')
        ])

    # Merge context + resume fields
    extra_resume_context = "\n".join([
        summary_text,
        experience_text,
        education_text
    ]).strip()

    combined_context = (context_sample + "\n" + extra_resume_context).strip()

    # ----------------------------------------------------------
    # RANDOM NOISE + VARIATION KEY
    # ----------------------------------------------------------
    style_noise_pool = [
        "use uncommon synonyms",
        "reorder ideas differently",
        "make phrasing more concise",
        "add creative wording twists",
        "slightly formal tone",
        "slightly casual tone",
        "shuffle activity order",
        "split compound ideas differently"
    ]
    random.shuffle(style_noise_pool)
    style_noise = ", ".join(style_noise_pool[:3])
    variation_key = f"{uuid.uuid4()}-{datetime.utcnow().timestamp()}"

    target_questions = [q for q in questions if q.get("parameter") in target_parameters]

    batch_prompt = (
        "You are an AI assistant generating short, career-related multiple-choice options.\n\n"
        "Generate options inspired by the user's Personal Interests, Hobbies, Exploration Interests, "
        "Motivation Drivers, Motivating Activities, and Achievements—without directly copying their context. "
        "Infer the user's underlying nature (e.g., creative, organized, exploratory) and tailor options accordingly. "
        "Ensure options remain relevant to the user's job title.\n\n"
        f"STRICT KNOWLEDGE BASE (use ONLY this content, no invention):\n{combined_context}\n\n"
        "QUESTIONS (JSON array):\n"
        f"{json.dumps([{'parameter': q['parameter'], 'question': q.get('question'), 'option_count': 2 * len([p.strip() for p in q['parameter'].split('+')])} for q in target_questions], indent=2)}\n\n"
        "Instructions:\n"
        "- For each question, generate EXACTLY the number of options in option_count.\n"
        "- Each option must rephrase or summarize ideas from the context.\n"
        "- DO NOT invent anything not in the context.\n"
        "- Keep options SHORT (2-5 words).\n"
        "- Start each option with a CAPITAL letter.\n"
        "- All options must be distinct.\n"
        f"- Apply variation rules: {style_noise}\n"
        f"- Variation key: {variation_key}\n\n"
        "Return ONLY valid JSON:\n"
        "{\"results\": [{\"parameter\": \"<exact parameter string>\", \"options\": [...]}, ...]}\n"
    )

    suggestions = []
    try:
        response = await _gemini_with_retry(batch_prompt)
        cleaned = clean_llm_json_response(response.text)
        parsed_batch = json.loads(cleaned)
        results_map = {r["parameter"]: r.get("options", []) for r in parsed_batch.get("results", [])}

        for q in target_questions:
            parameter = q.get("parameter")
            parameter_list = [p.strip() for p in parameter.split("+")]
            option_count = 2 * len(parameter_list)
            options = results_map.get(parameter, [])
            formatted = []
            for opt in options[:option_count]:
                formatted.append(opt.strip().lstrip("0123456789.- ").capitalize())
            while len(formatted) < option_count:
                formatted.append(f"Option {len(formatted) + 1}")
            suggestions.append({
                "parameter": parameter_list,
                "question": q.get("question"),
                "type": q.get("type"),
                "iconfilename": q.get("iconfilename"),
                "options": formatted,
            })
    except Exception as e:
        for q in target_questions:
            parameter = q.get("parameter")
            parameter_list = [p.strip() for p in parameter.split("+")]
            option_count = 2 * len(parameter_list)
            suggestions.append({
                "parameter": parameter_list,
                "question": q.get("question"),
                "type": q.get("type"),
                "iconfilename": q.get("iconfilename"),
                "options": [f"Option {i+1}" for i in range(option_count)],
                "error": str(e),
            })

    # ----------------------------------------------------------
    # SAVE RESULTS IN UPLOADS COLLECTION
    # ----------------------------------------------------------
    await db["uploads"].update_one(
        {"_id": latest_cv["_id"]},
        {"$set": {"anchor_questions_with_options": suggestions}}
    )

    return {"success": True, "suggestions": suggestions}


#--------------Second Flow--------------


async def generate_job_attribute_options_without_cv(user_id: str, db) -> dict:
    """
    Generates multiple-choice options for  Job attributes
    and saves them to the proceed_without_cv collection.
    """

    proceed_collection = db["proceed_without_cv"]
    answers_collection = db["answers_without_cv"]
    questions_collection = db["questions"]

    # Get latest proceed_without_cv doc
    latest_doc = await proceed_collection.find_one(
        {"user_id": user_id},
        sort=[("created_at", -1)]
    )
    if not latest_doc:
        raise HTTPException(status_code=404, detail="No proceed_without_cv doc found")

    latest_document_id = str(latest_doc["_id"])
    audience_type = latest_doc.get("audienceType")

    # Get missing answers
    missing_answers_cursor = answers_collection.find(
        {"user_id": user_id, "document_id": latest_document_id, "section": "Cv Missing"}
    )
    missing_answers = await missing_answers_cursor.to_list(length=None)

    if not missing_answers:
        raise HTTPException(status_code=404, detail="No missing answers found")

    # Build context string
    context_str = "\n".join(
        f"{a['parameter']}: {a.get('value')}" for a in missing_answers if a.get("value")
    )

    # Fetch job attributes questions
    questions_doc = await questions_collection.find_one({})
    if not questions_doc:
        raise HTTPException(status_code=404, detail="No questions document found")

    job_attributes = questions_doc.get("Job attributes", [])
    questions_for_user = [
        qa for qa in job_attributes if qa.get("audienceType") == audience_type
    ]

    if not questions_for_user:
        raise HTTPException(
            status_code=404, detail=f"No questions found for audienceType {audience_type}"
        )

    all_questions = [
        q for qa_group in questions_for_user for q in qa_group.get("questions", [])
    ]

    questions_meta = []
    for q in all_questions:
        parameter = q.get("parameter", "")
        parameter_list = [p.strip() for p in parameter.split("+")]
        option_count = 5 if len(parameter_list) == 1 else 2 * len(parameter_list)
        questions_meta.append({
            "parameter": parameter,
            "parameter_list": parameter_list,
            "question": q.get("question"),
            "type": q.get("type"),
            "iconfilename": q.get("iconfilename"),
            "limit": q.get("limit") or q.get("Limit"),
            "option_count": option_count,
        })

    batch_prompt = (
        "You are an AI assistant generating short, career-related multiple-choice options.\n\n"
        f"Audience Type: {audience_type}\n\n"
        f"Missing CV Context:\n{context_str}\n\n"
        "QUESTIONS (JSON array):\n"
        f"{json.dumps([{'parameter': m['parameter'], 'question': m['question'], 'option_count': m['option_count']} for m in questions_meta], indent=2)}\n\n"
        "Respond ONLY in valid JSON:\n"
        "{\"results\": [{\"parameter\": \"<exact parameter string>\", \"options\": [\"<short phrase>\", ...]}, ...]}\n\n"
        "RULES:\n"
        "- For each question, provide EXACTLY the number of options given in option_count.\n"
        "- Keep each option 2-5 words long.\n"
        "- Avoid numbering or letters (no A/B/C/... prefixes).\n"
        "- Make sure options fit the question meaningfully."
    )

    results = []
    try:
        response = await _gemini_with_retry(batch_prompt)
        cleaned = clean_llm_json_response(response.text)
        parsed_batch = json.loads(cleaned)
        results_map = {r["parameter"]: r.get("options", []) for r in parsed_batch.get("results", [])}

        for m in questions_meta:
            options = results_map.get(m["parameter"], [])
            formatted = [opt.strip() for opt in options[:m["option_count"]]]
            while len(formatted) < m["option_count"]:
                formatted.append(f"Option {len(formatted) + 1}")
            item = {
                "parameter": m["parameter_list"],
                "question": m["question"],
                "type": m["type"],
                "iconfilename": m["iconfilename"],
                "options": formatted,
            }
            if m["limit"] is not None:
                item["limit"] = m["limit"]
            results.append(item)
    except Exception as e:
        print(f"Batch generation failed, using fallback: {e}")
        for m in questions_meta:
            item = {
                "parameter": m["parameter_list"],
                "question": m["question"],
                "type": m["type"],
                "iconfilename": m["iconfilename"],
                "options": [f"Option {i+1}" for i in range(m["option_count"])],
            }
            if m["limit"] is not None:
                item["limit"] = m["limit"]
            results.append(item)

    # Save results
    await proceed_collection.update_one(
        {"_id": latest_doc["_id"]},
        {"$set": {"job_questions_with_options": results, "updated_at": datetime.utcnow()}}
    )

    return {
        "success": True,
        "latest_document_id": latest_document_id,
        "audience_type": audience_type,
        "suggestions": results,
    }


async def generate_anchor_options_from_answers_without_cv(
    user_id: str, questions, model, get_database
) -> Dict[str, Any]:
    """
    Generate multiple-choice options for Anchor attributes based ONLY on:
    - Personal Interests + Hobbies + Exploration Interest + Motivation Drivers + Motivating Activities
    - Achievements (text field only)
    Fetch answers from answers_without_cv based on the latest proceed_without_cv.document_id,
    generate options using LLM, and save them back into proceed_without_cv.
    """

    OPTIONS_PER_PARAMETER = 2

    target_parameters = {
        "Creative Inclinations + Organizational Skills + Competency + Personality Traits",
        "Newly Acquired Skills + Emerging Tech Awareness + Future Study Intent",
    }

    db = await get_database()
    latest_proceed = await db["proceed_without_cv"].find(
        {"user_id": user_id}
    ).sort("created_at", -1).to_list(length=1)

    if not latest_proceed:
        raise HTTPException(status_code=404, detail="No proceed_without_cv found for user")

    document = latest_proceed[0]
    document_id = document["_id"]
    audience_type = document.get("audienceType", "General")

    answers_cursor = db["answers_without_cv"].find(
        {
            "user_id": user_id,
            "document_id": str(document_id),
            "parameter": {
                "$in": [
                    "Personal Interests + Hobbies + Exploration Interest + Motivation Drivers + Motivating Activities",
                    "Achievements",
                ]
            },
        }
    )
    answers = await answers_cursor.to_list(length=None)

    if not answers:
        raise HTTPException(status_code=404, detail="No required anchor answers found")

    base_free_text_map = {}
    for ans in answers:
        param = ans["parameter"]
        val = ans["value"]
        if isinstance(val, str):
            base_free_text_map[param] = val
        elif isinstance(val, dict) and "text" in val:
            base_free_text_map[param] = val["text"]

    non_empty_texts = [t for t in base_free_text_map.values() if t]
    random.shuffle(non_empty_texts)
    context_sample = "\n".join(non_empty_texts)

    variation_key = f"{uuid.uuid4()}-{datetime.utcnow().timestamp()}"
    style_noise_pool = [
        "use uncommon synonyms",
        "reorder ideas differently",
        "make phrasing more concise",
        "add creative wording twists",
        "slightly formal tone",
        "slightly casual tone",
        "shuffle activity order",
        "split compound ideas differently",
    ]
    random.shuffle(style_noise_pool)
    style_noise = ", ".join(style_noise_pool[:3])

    target_questions = [q for q in questions if q.get("parameter") in target_parameters]

    batch_prompt = (
        "You are an AI assistant generating multiple-choice options for career-related questions.\n\n"
        f"STRICT KNOWLEDGE BASE (rephrase ONLY from this, do not add new ideas):\n{context_sample}\n\n"
        "QUESTIONS (JSON array):\n"
        f"{json.dumps([{'parameter': q['parameter'], 'question': q.get('question'), 'option_count': OPTIONS_PER_PARAMETER * len([p.strip() for p in q['parameter'].split('+')])} for q in target_questions], indent=2)}\n\n"
        "Instructions:\n"
        "- For each question, generate EXACTLY the number of options in option_count.\n"
        "- Each option must be a direct rephrasing, splitting, or summarizing of the free-text answers.\n"
        "- DO NOT invent anything that is not explicitly present in the free-text answers.\n"
        "- Keep each option SHORT (2-5 words).\n"
        "- Ensure each execution produces DIFFERENT wording.\n"
        f"- Apply variation rules: {style_noise}\n"
        f"- Variation key: {variation_key}\n\n"
        "Return ONLY valid JSON:\n"
        "{\"results\": [{\"parameter\": \"<exact parameter string>\", \"options\": [...]}, ...]}\n"
    )

    suggestions = []
    try:
        response = await _gemini_with_retry(batch_prompt)
        cleaned = clean_llm_json_response(response.text)
        parsed_batch = json.loads(cleaned)
        results_map = {r["parameter"]: r.get("options", []) for r in parsed_batch.get("results", [])}

        for q in target_questions:
            parameter = q.get("parameter")
            parameter_list = [p.strip() for p in parameter.split("+")]
            option_count = OPTIONS_PER_PARAMETER * len(parameter_list)
            options = results_map.get(parameter, [])
            formatted = [
                options[i].strip() if i < len(options) else f"Option {i+1}"
                for i in range(option_count)
            ]
            suggestions.append({
                "parameter": parameter_list,
                "question": q.get("question"),
                "type": q.get("type"),
                "iconfilename": q.get("iconfilename"),
                "options": formatted,
            })
    except Exception as e:
        for q in target_questions:
            parameter = q.get("parameter")
            parameter_list = [p.strip() for p in parameter.split("+")]
            option_count = OPTIONS_PER_PARAMETER * len(parameter_list)
            suggestions.append({
                "parameter": parameter_list,
                "question": q.get("question"),
                "type": q.get("type"),
                "iconfilename": q.get("iconfilename"),
                "options": [f"Option {i+1}" for i in range(option_count)],
                "error": str(e),
            })
    await db["proceed_without_cv"].update_one(
        {"_id": document_id},
        {"$set": {"anchor_questions_with_options": suggestions}},
    )

    return {
        "success": True,
        "document_id": str(document_id),
        "audienceType": audience_type,
        "suggestions": suggestions,
    }
